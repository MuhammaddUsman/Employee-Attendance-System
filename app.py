from flask import Flask, render_template_string, request, jsonify, redirect, url_for, send_file, abort, session
from datetime import datetime, timedelta
import json
import os
import csv
import re
import io
import time
import hmac
import urllib.request
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
from zoneinfo import ZoneInfo

app = Flask(__name__)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=bool(os.environ.get('VERCEL')),
)

app.secret_key = os.environ.get('SECRET_KEY', 'dev_secret_change_me')

ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'admin')
TIMEZONE = ZoneInfo(os.environ.get('TIMEZONE', 'Asia/Karachi'))

# Storage: Upstash Redis (persistent, used on Vercel) if configured,
# otherwise a local JSON file (fine for local development).
REDIS_URL = os.environ.get('KV_REST_API_URL') or os.environ.get('UPSTASH_REDIS_REST_URL')
REDIS_TOKEN = os.environ.get('KV_REST_API_TOKEN') or os.environ.get('UPSTASH_REDIS_REST_TOKEN')
REDIS_KEY = 'attendance_data'
DATA_FILE = os.path.join('/tmp' if os.environ.get('VERCEL') else '.', 'attendance_data.json')


P_EMP, P_ATT = 'att:employees', 'att:attendance'
K_CNT, K_MIG = 'att:emp_counter', 'att:migrated'
_migrated = False


def now_local():
    return datetime.now(TIMEZONE)


def _use_redis():
    return bool(REDIS_URL and REDIS_TOKEN)


def _cmd(*args):
    req = urllib.request.Request(
        REDIS_URL.rstrip('/'),
        data=json.dumps(list(args)).encode('utf-8'),
        headers={'Authorization': f'Bearer {REDIS_TOKEN}', 'Content-Type': 'application/json'},
        method='POST',
    )
    with urllib.request.urlopen(req, timeout=8) as resp:
        out = json.loads(resp.read().decode('utf-8'))
    if isinstance(out, dict) and out.get('error'):
        raise RuntimeError(out['error'])
    return out.get('result')


def _emp_num(eid):
    m = re.fullmatch(r'EMP(\d+)', eid or '')
    return int(m.group(1)) if m else 0


# ---- local JSON file (development only) ----
def _local_load():
    d = {}
    if os.path.exists(DATA_FILE):
        try:
            with open(DATA_FILE, 'r') as f:
                d = json.load(f)
        except Exception:
            d = {}
    if isinstance(d.get('attendance'), list) or (isinstance(d.get('employees'), dict) and P_EMP not in d):
        old, d = d, {P_EMP: {}, P_ATT: {}, K_CNT: {}}   # convert the old single-blob format
        for eid, e in (old.get('employees') or {}).items():
            d[P_EMP][eid] = json.dumps(e)
        for r in old.get('attendance') or []:
            d[P_ATT].setdefault(f"{r['date']}:{r['employee_id']}", json.dumps(r))
        d[K_CNT]['v'] = str(max([_emp_num(e) for e in d[P_EMP]] or [0]))
    for k in (P_EMP, P_ATT, K_CNT):
        d.setdefault(k, {})
    return d


def _local_save(d):
    with open(DATA_FILE, 'w') as f:
        json.dump(d, f, indent=1)


# ---- one-time move from the old single-key storage to per-record storage ----
def ensure_migrated():
    global _migrated
    if _migrated or not _use_redis():
        return
    if _cmd('GET', K_MIG) is None:
        raw = _cmd('GET', REDIS_KEY)           # the old blob is kept as a backup
        if raw:
            old = json.loads(raw)
            emps = old.get('employees') or {}
            if emps:
                args = []
                for eid, e in emps.items():
                    args += [eid, json.dumps(e)]
                _cmd('HSET', P_EMP, *args)
            recs = {}
            for r in old.get('attendance') or []:
                recs.setdefault(f"{r['date']}:{r['employee_id']}", json.dumps(r))
            if recs:
                args = []
                for k, v in recs.items():
                    args += [k, v]
                _cmd('HSET', P_ATT, *args)
            _cmd('SETNX', K_CNT, str(max([_emp_num(e) for e in emps] or [0])))
        _cmd('SET', K_MIG, '1')
    _migrated = True


# ---- hash primitives (same behaviour on Redis and on the local file) ----
def h_getall(name):
    if _use_redis():
        ensure_migrated()
        res = _cmd('HGETALL', name) or []
        if isinstance(res, dict):
            return res
        return {res[i]: res[i + 1] for i in range(0, len(res), 2)}
    return dict(_local_load()[name])


def h_get(name, field):
    if _use_redis():
        ensure_migrated()
        return _cmd('HGET', name, field)
    return _local_load()[name].get(field)


def h_set(name, field, value):
    if _use_redis():
        ensure_migrated()
        _cmd('HSET', name, field, value)
        return
    d = _local_load()
    d[name][field] = value
    _local_save(d)


def h_setnx(name, field, value):
    # atomic "create only if it does not exist"
    if _use_redis():
        ensure_migrated()
        return _cmd('HSETNX', name, field, value) == 1
    d = _local_load()
    if field in d[name]:
        return False
    d[name][field] = value
    _local_save(d)
    return True


def h_del(name, *fields):
    if not fields:
        return
    if _use_redis():
        ensure_migrated()
        _cmd('HDEL', name, *fields)
        return
    d = _local_load()
    for f in fields:
        d[name].pop(f, None)
    _local_save(d)


def next_counter():
    if _use_redis():
        ensure_migrated()
        return int(_cmd('INCR', K_CNT))
    d = _local_load()
    n = int(d[K_CNT].get('v', '0')) + 1
    d[K_CNT]['v'] = str(n)
    _local_save(d)
    return n


# ---- employees and records ----
def rec_key(date, emp_id):
    return f'{date}:{emp_id}'


def store_get_employee(eid):
    v = h_get(P_EMP, eid)
    return json.loads(v) if v else None


def store_set_employee(eid, emp):
    h_set(P_EMP, eid, json.dumps(emp))


def store_create_employee(eid, emp):
    return h_setnx(P_EMP, eid, json.dumps(emp))


def store_delete_employee(eid):
    h_del(P_EMP, eid)
    h_del(P_ATT, *[k for k in h_getall(P_ATT) if k.endswith(':' + eid)])


def store_get_record(date, eid):
    v = h_get(P_ATT, rec_key(date, eid))
    return json.loads(v) if v else None


def store_set_record(rec):
    h_set(P_ATT, rec_key(rec['date'], rec['employee_id']), json.dumps(rec))


def store_create_record(rec):
    return h_setnx(P_ATT, rec_key(rec['date'], rec['employee_id']), json.dumps(rec))


def store_delete_record(date, eid):
    h_del(P_ATT, rec_key(date, eid))


def load_data():
    employees = {k: json.loads(v) for k, v in h_getall(P_EMP).items()}
    attendance = [json.loads(v) for v in h_getall(P_ATT).values()]
    attendance.sort(key=lambda r: (r['date'], r.get('clock_in') or ''))
    return {'employees': employees, 'attendance': attendance}


BASE_TEMPLATE = r'''
<!DOCTYPE html>
<html>
<head>
    <title>{{ company_name }} - Attendance</title>
    <style>
        :root{--navy:#0f2547;--blue:#2563eb;--bg:#f1f5f9;--text:#0f172a;--muted:#64748b;--line:#e2e8f0;--green:#15803d;--red:#b91c1c}
        *{margin:0;padding:0;box-sizing:border-box}
        body{font-family:'Inter','Segoe UI',system-ui,-apple-system,Roboto,Arial,sans-serif;background:var(--bg);color:var(--text);font-size:15px;line-height:1.5}
        .topbar{background:var(--navy);position:sticky;top:0;z-index:10;box-shadow:0 1px 4px rgba(0,0,0,.25)}
        .topbar-inner{max-width:1200px;margin:0 auto;padding:0 24px;display:flex;align-items:center;justify-content:space-between;gap:16px;flex-wrap:wrap;min-height:60px}
        .brand{color:#fff;font-weight:700;font-size:1.1rem;display:flex;align-items:center;gap:10px}
        .brand-mark{width:32px;height:32px;border-radius:8px;background:var(--blue);display:inline-flex;align-items:center;justify-content:center;font-size:1rem}
        .nav{display:flex;gap:4px;flex-wrap:wrap;padding:8px 0}
        .nav a{color:#cbd5e1;text-decoration:none;padding:8px 14px;border-radius:8px;font-weight:500;font-size:.92rem;transition:background .15s,color .15s}
        .nav a:hover{background:rgba(255,255,255,.1);color:#fff}
        .nav a.active{background:rgba(255,255,255,.18);color:#fff}
        .container{max-width:1200px;margin:0 auto;padding:28px 24px 40px}
        .footer{text-align:center;color:var(--muted);font-size:.82rem;padding:0 24px 28px}
        .page-title{font-size:1.25rem;font-weight:700;color:var(--text);margin-bottom:16px}
        .card{background:#fff;border:1px solid var(--line);border-radius:12px;padding:24px;margin-bottom:20px;box-shadow:0 1px 2px rgba(15,23,42,.05);overflow-x:auto}
        .stats-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:16px;margin-bottom:20px}
        .stat-card{background:#fff;border:1px solid var(--line);border-left:4px solid var(--blue);border-radius:12px;padding:16px 20px}
        .stat-card .num{font-size:1.9rem;font-weight:700;line-height:1.1}
        .stat-card .lbl{color:var(--muted);font-size:.78rem;margin-top:4px;text-transform:uppercase;letter-spacing:.06em}
        .form-group{margin-bottom:18px}
        .form-group label,.card label{display:block;margin-bottom:6px;font-weight:600;color:#334155;font-size:.9rem}
        .form-group input,.form-group select,.card input[type=date],.card input[type=month]{width:100%;padding:10px 12px;border:1px solid #cbd5e1;border-radius:8px;font-size:.95rem;background:#fff;font-family:inherit}
        .card input[type=date],.card input[type=month]{width:auto;display:inline-block}
        .card form label{display:inline-block;margin-bottom:0}
        input:focus,select:focus{outline:none;border-color:var(--blue);box-shadow:0 0 0 3px rgba(37,99,235,.15)}
        .btn{display:inline-block;padding:10px 20px;background:var(--blue);color:#fff;border:none;border-radius:8px;cursor:pointer;font-size:.92rem;font-weight:600;text-decoration:none;font-family:inherit;transition:background .15s}
        .btn:hover{background:#1d4ed8}
        .btn-success{background:var(--green)}.btn-success:hover{background:#166534}
        .btn-danger{background:var(--red)}.btn-danger:hover{background:#991b1b}
        table{width:100%;border-collapse:collapse;margin-top:8px;min-width:560px}
        table th{background:#f8fafc;color:var(--muted);padding:12px 14px;text-align:left;font-weight:600;text-transform:uppercase;font-size:.74rem;letter-spacing:.05em;border-bottom:2px solid var(--line)}
        table td{padding:12px 14px;border-bottom:1px solid var(--line);color:#1e293b}
        table tbody tr:nth-child(even){background:#fafbfd}
        table tbody tr:hover{background:#f1f5f9}
        .status-in,.status-out{font-weight:600;font-size:.8rem;padding:4px 12px;border-radius:999px;display:inline-block}
        .status-in{color:#166534;background:#dcfce7}
        .status-out{color:#475569;background:#e2e8f0}
        .clock-display{text-align:center;font-size:3rem;font-weight:700;color:var(--navy);margin:6px 0 20px;font-variant-numeric:tabular-nums}
        .date-display{text-align:center;color:var(--muted);font-weight:500;margin-bottom:4px}
        .success-msg,.error-msg{padding:12px 16px;border-radius:8px;margin-bottom:18px;font-weight:500;font-size:.92rem}
        .success-msg{background:#f0fdf4;color:#166534;border:1px solid #bbf7d0}
        .error-msg{background:#fef2f2;color:#991b1b;border:1px solid #fecaca}
        .button-group{display:flex;gap:12px;flex-wrap:wrap}
        .button-group .btn{flex:1;min-width:150px;padding:12px 20px}
        @media(max-width:640px){.container{padding:18px 14px 30px}.card{padding:16px}.clock-display{font-size:2.3rem}.topbar-inner{padding:0 14px}}
    </style>
</head>
<body>
    {% set p = request.path %}
    <div class="topbar"><div class="topbar-inner">
        <div class="brand"><span class="brand-mark">{{ company_name[:1] }}</span>{{ company_name }}</div>
        <div class="nav">
            <a href="/" class="{{ 'active' if p == '/' }}">Clock In/Out</a>
            <a href="/employees" class="{{ 'active' if p == '/employees' }}">Employees</a>
            {% if is_admin %}
            <a href="/records" class="{{ 'active' if p == '/records' }}">Attendance Records</a>
            <a href="/reports" class="{{ 'active' if p == '/reports' }}">Reports</a>
            <a href="/export">Export Excel</a>
            {% endif %}
            {% if is_admin %}<a href="/logout">Logout</a>{% else %}<a href="/login" class="{{ 'active' if p == '/login' }}">Admin Login</a>{% endif %}
        </div>
    </div></div>
    <div class="container">
        {% block content %}{% endblock %}
    </div>
    <div class="footer">&copy; {{ company_name }} &middot; Attendance Management System</div>
    <script>
        function updateClock() {
            const now = new Date();
            const timeStr = now.toLocaleTimeString('en-US', { hour: '2-digit', minute: '2-digit', second: '2-digit' });
            const dateStr = now.toLocaleDateString('en-US', { weekday: 'long', year: 'numeric', month: 'long', day: 'numeric' });
            const clockEl = document.getElementById('live-clock');
            const dateEl = document.getElementById('live-date');
            if (clockEl) {
                clockEl.textContent = timeStr;
            }
            if (dateEl) {
                dateEl.textContent = dateStr;
            }
        }
        setInterval(updateClock, 1000);
        updateClock();
    </script>
</body>
</html>
'''

LOGIN_TEMPLATE = BASE_TEMPLATE.replace('{% block content %}{% endblock %}', '''
    <div class="card" style="max-width:480px;margin:0 auto;">
        <h2 class="page-title">Admin Login</h2>
        {% if message %}
            <div class="error-msg">{{ message }}</div>
        {% endif %}
        <form method="POST" action="/login">
            <input type="hidden" name="next" value="{{ next or '' }}">
            <div class="form-group">
                <label>Admin Password:</label>
                <input type="password" name="password" required placeholder="Enter admin password">
            </div>
            <div style="display:flex;gap:10px;">
                <button type="submit" class="btn btn-success">Login</button>
                <a href="/" class="btn">Back</a>
            </div>
        </form>
    </div>
''')


HOME_TEMPLATE = BASE_TEMPLATE.replace('{% block content %}{% endblock %}', '''
    <div class="stats-grid">
        <div class="stat-card"><div class="num">{{ total_employees }}</div><div class="lbl">Total Employees</div></div>
        <div class="stat-card"><div class="num">{{ working_now }}</div><div class="lbl">Working Now</div></div>
        <div class="stat-card"><div class="num">{{ completed_today }}</div><div class="lbl">Completed Today</div></div>
    </div>
    <div class="card">
        <h2 class="page-title">Clock In/Out</h2>
        <div class="date-display" id="live-date"></div>
        <div class="clock-display" id="live-clock"></div>
        {% if message %}
        <div class="{% if success %}success-msg{% else %}error-msg{% endif %}">
            {{ message }}
        </div>
        {% endif %}
        <form method="POST" action="/clock">
            <div class="form-group">
                <label>Select Employee:</label>
                <select name="employee_id" required>
                    <option value="">-- Choose Employee --</option>
                    {% for emp_id, emp in employees.items() %}
                    <option value="{{ emp_id }}">{{ emp.name }} ({{ emp.department }})</option>
                    {% endfor %}
                </select>
            </div>
            <div class="form-group">
                <label>Your PIN:</label>
                <input type="password" name="pin" inputmode="numeric" pattern="[0-9]{4,6}" maxlength="6" autocomplete="off" required placeholder="Enter your PIN">
            </div>
            <div class="button-group">
                <button type="submit" name="action" value="in" class="btn btn-success">Clock In</button>
                <button type="submit" name="action" value="out" class="btn btn-danger">Clock Out</button>
            </div>
        </form>
    </div>
    
    <div class="card">
        <h3 class="page-title">Today's Activity</h3>
        <table>
            <thead>
                <tr>
                    <th>Employee</th>
                    <th>Department</th>
                    <th>Clock In</th>
                    <th>Clock Out</th>
                    <th>Status</th>
                </tr>
            </thead>
            <tbody>
                {% for record in today_records %}
                <tr>
                    <td><strong>{{ record.employee_name }}</strong></td>
                    <td>{{ record.department }}</td>
                    <td>{{ record.clock_in }}</td>
                    <td>{{ record.clock_out or '-' }}</td>
                    <td>
                        <span class="{% if record.clock_out %}status-out{% else %}status-in{% endif %}">
                            {% if record.clock_out %}Clocked Out{% else %}Working{% endif %}
                        </span>
                    </td>
                </tr>
                {% endfor %}
            </tbody>
        </table>
    </div>
''')


EMPLOYEES_TEMPLATE = BASE_TEMPLATE.replace('{% block content %}{% endblock %}', '''
    <div class="card">
        <h2 class="page-title">Add New Employee</h2>
        {% if message %}
        <div class="success-msg">{{ message }}</div>
        {% endif %}
        {% if is_admin %}
        <form method="POST" action="/add_employee">
            <div class="form-group">
                <label>Employee Name:</label>
                <input type="text" name="name" required placeholder="Enter full name">
            </div>
            <div class="form-group">
                <label>Department:</label>
                <input type="text" name="department" required placeholder="e.g., IT, HR, Sales">
            </div>
            <div class="form-group">
                <label>Position:</label>
                <input type="text" name="position" required placeholder="e.g., Manager, Developer">
            </div>
            <div class="form-group">
                <label>Employee PIN (4-6 digits):</label>
                <input type="password" name="pin" inputmode="numeric" pattern="[0-9]{4,6}" maxlength="6" required autocomplete="new-password" placeholder="e.g., 4821">
            </div>
            <button type="submit" class="btn">Add Employee</button>
        </form>
        {% else %}
        <p style="color:#64748b;">Only an admin can add employees. <a href="/login?next=/employees">Admin login</a></p>
        {% endif %}
    </div>
    
    <div class="card">
        <h3 class="page-title">Employee List</h3>
        <table>
            <thead>
                <tr>
                    <th>ID</th>
                    <th>Name</th>
                    <th>Department</th>
                    <th>Position</th>
                    <th>Join Date</th>
                    <th>Actions</th>
                </tr>
            </thead>
            <tbody>
                {% for emp_id, emp in employees.items() %}
                <tr>
                    <td><strong>{{ emp_id }}</strong></td>
                    <td>{{ emp.name }}{% if is_admin and not emp.pin_hash %} <span class="status-out" style="color:#b91c1c;background:#fee2e2;">No PIN</span>{% endif %}</td>
                    <td>{{ emp.department }}</td>
                    <td>{{ emp.position }}</td>
                    <td>{{ emp.join_date }}</td>
                    <td>
                        {% if is_admin %}
                            <form method="POST" action="/set_pin" style="display:inline-flex;gap:6px;align-items:center;">
                                <input type="hidden" name="emp_id" value="{{ emp_id }}">
                                <input type="password" name="pin" inputmode="numeric" pattern="[0-9]{4,6}" maxlength="6" required autocomplete="new-password" placeholder="New PIN" style="width:100px;padding:7px 9px;border:1px solid #cbd5e1;border-radius:8px;">
                                <button type="submit" class="btn" style="padding:7px 12px;">Set PIN</button>
                            </form>
                            <form method="POST" action="/delete_employee" style="display:inline;" onsubmit="return confirm('Delete this employee and all of their attendance records?');">
                                <input type="hidden" name="emp_id" value="{{ emp_id }}">
                                <button type="submit" class="btn btn-danger" style="margin-left:6px;padding:7px 12px;">Delete</button>
                            </form>
                        {% else %}
                            <!-- Non-admin: no delete option -->
                            <span style="color:#718096;font-weight:600;">—</span>
                        {% endif %}
                    </td>
                </tr>
                {% endfor %}
            </tbody>
        </table>
    </div>
''')


RECORDS_TEMPLATE = BASE_TEMPLATE.replace('{% block content %}{% endblock %}', '''
    <div class="card">
        <h2 class="page-title">Attendance Records</h2>
        {% if message %}<div class="success-msg">{{ message }}</div>{% endif %}
        <p style="margin-bottom:16px;"><a href="/edit_record" class="btn">+ Add / Correct Record</a></p>
        <form method="GET" action="/records" style="margin-bottom:16px;">
            <label>Filter by date:</label>
            <input type="date" name="date" value="{{ filter_date or '' }}" style="padding:8px;border-radius:8px;border:1px solid #e2e8f0;margin-left:8px;">
            <button class="btn" style="margin-left:8px;">Apply</button>
            <a href="/records" class="btn" style="margin-left:8px;">Clear</a>
        </form>
        <form method="GET" action="/export" style="margin-bottom:16px;">
            <label>Download monthly Excel sheet:</label>
            <input type="month" name="month" value="{{ current_month }}" style="padding:8px;border-radius:8px;border:1px solid #e2e8f0;margin-left:8px;">
            <button class="btn" style="margin-left:8px;">Download</button>
        </form>
        <table>
            <thead>
                <tr>
                    <th>Date</th>
                    <th>Employee</th>
                    <th>Department</th>
                    <th>Clock In</th>
                    <th>Clock Out</th>
                    <th>Time Worked</th>
                    <th>Actions</th>
                </tr>
            </thead>
            <tbody>
                {% for record in records %}
                <tr>
                    <td>{{ record.date }}</td>
                    <td>{{ record.employee_name }}</td>
                    <td>{{ record.department }}</td>
                    <td>{{ record.clock_in }}</td>
                    <td>{{ record.clock_out or '-' }}</td>
                    <td>{{ record.hours_worked }}</td>
                    <td>
                        <a href="/edit_record?date={{ record.date }}&emp_id={{ record.employee_id }}" class="btn" style="padding:6px 12px;">Edit</a>
                        <form method="POST" action="/delete_record" style="display:inline;" onsubmit="return confirm('Delete this record?');">
                            <input type="hidden" name="date" value="{{ record.date }}">
                            <input type="hidden" name="emp_id" value="{{ record.employee_id }}">
                            <button type="submit" class="btn btn-danger" style="padding:6px 12px;">Delete</button>
                        </form>
                    </td>
                </tr>
                {% endfor %}
            </tbody>
        </table>
    </div>
''')


EDIT_TEMPLATE = BASE_TEMPLATE.replace('{% block content %}{% endblock %}', '''
    <div class="card" style="max-width:560px;margin:0 auto;">
        <h2 class="page-title">{{ 'Edit Record' if editing else 'Add / Correct Record' }}</h2>
        {% if message %}<div class="error-msg">{{ message }}</div>{% endif %}
        <form method="POST" action="/save_record">
            <div class="form-group">
                <label>Employee:</label>
                {% if editing %}
                <input type="hidden" name="emp_id" value="{{ emp_id }}">
                <input type="text" value="{{ employees[emp_id].name if emp_id in employees else emp_id }}" disabled>
                {% else %}
                <select name="emp_id" required>
                    <option value="">-- Choose Employee --</option>
                    {% for eid, e in employees.items() %}
                    <option value="{{ eid }}" {{ 'selected' if eid == emp_id }}>{{ e.name }} ({{ e.department }})</option>
                    {% endfor %}
                </select>
                {% endif %}
            </div>
            <div class="form-group">
                <label>Date:</label>
                {% if editing %}
                <input type="hidden" name="date" value="{{ date }}">
                <input type="text" value="{{ date }}" disabled>
                {% else %}
                <input type="date" name="date" value="{{ date }}" required>
                {% endif %}
            </div>
            <div class="form-group">
                <label>Clock In:</label>
                <input type="time" step="1" name="clock_in" value="{{ clock_in }}" required>
            </div>
            <div class="form-group">
                <label>Clock Out (leave empty if still working):</label>
                <input type="time" step="1" name="clock_out" value="{{ clock_out }}">
            </div>
            <div style="display:flex;gap:10px;">
                <button type="submit" class="btn btn-success">Save</button>
                <a href="/records" class="btn" style="background:#64748b;">Cancel</a>
            </div>
        </form>
    </div>
''')


REPORTS_TEMPLATE = BASE_TEMPLATE.replace('{% block content %}{% endblock %}', '''
    <div class="card">
        <h2 class="page-title">Weekly Attendance Report</h2>
        <table>
            <thead>
                <tr>
                    <th>Employee</th>
                    <th>Department</th>
                    <th>Days Present</th>
                    <th>Total Time</th>
                    <th>Average Time/Day</th>
                </tr>
            </thead>
            <tbody>
                {% for emp_id, report in weekly_report.items() %}
                <tr>
                    <td>{{ report.name }}</td>
                    <td>{{ report.department }}</td>
                    <td>{{ report.days_present }}</td>
                    <td>{{ report.total_text }}</td>
                    <td>{{ report.avg_text }}</td>
                </tr>
                {% endfor %}
            </tbody>
        </table>
    </div>
''')


def hours_between(date_str, clock_in, clock_out):
    """Hours worked on a record. Returns None if the times are invalid
    (for example clock-out earlier than clock-in)."""
    try:
        a = datetime.strptime(f"{date_str} {clock_in}", '%Y-%m-%d %H:%M:%S')
        b = datetime.strptime(f"{date_str} {clock_out}", '%Y-%m-%d %H:%M:%S')
    except Exception:
        return None
    hours = (b - a).total_seconds() / 3600
    if hours < 0:
        return None
    return hours


def format_duration(hours, signed=False):
    """7.5 -> '7h 30m'. Under one minute -> '36s'. signed=True adds +/-."""
    if hours is None:
        return ''
    secs = round(abs(hours) * 3600)
    if 0 < secs < 60 and not signed:
        return f'{secs}s'
    h, m = divmod(round(secs / 60), 60)
    text = f'{h}h {m:02d}m'
    if signed:
        text = ('-' if hours < 0 else '+') + text if (h or m) else text
    return text


def parse_bool_param(val):
    if val is None:
        return None
    return str(val) in ('1', 'true', 'True', 'yes', 'Yes')


def parse_time(text):
    text = (text or '').strip()
    for fmt in ('%H:%M:%S', '%H:%M'):
        try:
            return datetime.strptime(text, fmt).strftime('%H:%M:%S')
        except ValueError:
            pass
    return None


def valid_date(text):
    try:
        datetime.strptime(text or '', '%Y-%m-%d')
        return True
    except ValueError:
        return False


def valid_pin(pin):
    return bool(re.fullmatch(r'[0-9]{4,6}', pin or ''))


def safe_next(url):
    if url and url.startswith('/') and not url.startswith('//') and '\\' not in url:
        return url
    return None


def admin_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get('is_admin'):
            if request.method == 'GET':
                target = request.full_path if request.query_string else request.path
            else:
                target = url_for('employees')
            return redirect(url_for('login', next=target))
        return f(*args, **kwargs)
    return wrapper


@app.context_processor
def inject_company():
    return {'company_name': COMPANY_NAME}


@app.route('/login', methods=['GET', 'POST'])
def login():
    nxt = safe_next(request.values.get('next'))
    if request.method == 'POST':
        pwd = request.form.get('password', '')
        if hmac.compare_digest(pwd.encode('utf-8'), ADMIN_PASSWORD.encode('utf-8')):
            session['is_admin'] = True
            return redirect(nxt or url_for('employees'))
        time.sleep(1)  # slows down password guessing
        return render_template_string(LOGIN_TEMPLATE, message='Invalid password', next=nxt, is_admin=False)
    return render_template_string(LOGIN_TEMPLATE, next=nxt, is_admin=session.get('is_admin', False))


@app.route('/logout')
def logout():
    session.pop('is_admin', None)
    return redirect(url_for('index'))


@app.route('/')
def index():
    data = load_data()
    today = now_local().strftime('%Y-%m-%d')
    today_records = [r for r in data['attendance'] if r['date'] == today]
    today_records = sorted(today_records, key=lambda r: r.get('clock_in', ''), reverse=False)
    message = request.args.get('message')
    success = parse_bool_param(request.args.get('success'))
    return render_template_string(HOME_TEMPLATE,
                                 employees=data['employees'],
                                 today_records=today_records,
                                 message=message,
                                 success=success,
                                 total_employees=len(data['employees']),
                                 working_now=sum(1 for r in today_records if not r.get('clock_out')),
                                 completed_today=sum(1 for r in today_records if r.get('clock_out')),
                                 is_admin=session.get('is_admin', False))


MAX_PIN_FAILS = 5
LOCK_SECONDS = 300


def check_pin(emp_id, emp, pin):
    # returns an error message, or None when the PIN is correct
    if not emp.get('pin_hash'):
        return 'No PIN is set for this employee. Please ask the admin to set one.'
    now_ts = time.time()
    if emp.get('locked_until', 0) > now_ts:
        mins = int((emp['locked_until'] - now_ts) // 60) + 1
        return f'Too many wrong attempts. Try again in {mins} minute(s).'
    if not pin or not check_password_hash(emp['pin_hash'], pin):
        emp['fails'] = emp.get('fails', 0) + 1
        if emp['fails'] >= MAX_PIN_FAILS:
            emp['fails'] = 0
            emp['locked_until'] = now_ts + LOCK_SECONDS
            store_set_employee(emp_id, emp)
            return 'Too many wrong attempts. This employee is locked for 5 minutes.'
        store_set_employee(emp_id, emp)
        return f'Incorrect PIN. {MAX_PIN_FAILS - emp["fails"]} attempt(s) left.'
    if emp.get('fails') or emp.get('locked_until'):
        emp['fails'] = 0
        emp['locked_until'] = 0
        store_set_employee(emp_id, emp)
    return None


@app.route('/clock', methods=['POST'])
def clock():
    employee_id = request.form.get('employee_id', '')
    action = request.form.get('action')
    pin = request.form.get('pin', '').strip()

    emp = store_get_employee(employee_id) if employee_id else None
    if not emp:
        return redirect(url_for('index', message='Employee not found', success=0))
    err = check_pin(employee_id, emp, pin)
    if err:
        return redirect(url_for('index', message=err, success=0))

    now = now_local()
    today = now.strftime('%Y-%m-%d')

    if action == 'in':
        rec = {
            'employee_id': employee_id,
            'employee_name': emp['name'],
            'department': emp['department'],
            'date': today,
            'clock_in': now.strftime('%H:%M:%S'),
            'clock_out': None,
        }
        if store_create_record(rec):   # atomic: a second click cannot create a duplicate
            return redirect(url_for('index', message=f'{emp["name"]} clocked in successfully!', success=1))
        existing = store_get_record(today, employee_id) or {}
        msg = 'Already clocked out for today.' if existing.get('clock_out') else 'Already clocked in!'
        return redirect(url_for('index', message=msg, success=0))

    if action == 'out':
        rec = store_get_record(today, employee_id)
        if not rec:
            return redirect(url_for('index', message='No clock-in record found!', success=0))
        if rec.get('clock_out'):
            return redirect(url_for('index', message='Already clocked out!', success=0))
        rec['clock_out'] = now.strftime('%H:%M:%S')
        store_set_record(rec)
        return redirect(url_for('index', message=f'{emp["name"]} clocked out successfully!', success=1))

    return redirect(url_for('index', message='Invalid action', success=0))


@app.route('/employees')
def employees():
    data = load_data()
    message = request.args.get('message')
    return render_template_string(EMPLOYEES_TEMPLATE, employees=data['employees'], message=message, is_admin=session.get('is_admin', False))


@app.route('/add_employee', methods=['POST'])
@admin_required
def add_employee():
    name = request.form.get('name', '').strip()
    department = request.form.get('department', '').strip()
    position = request.form.get('position', '').strip()
    pin = request.form.get('pin', '').strip()

    if not name or not department or not position:
        return redirect(url_for('employees', message='All fields are required'))
    if not valid_pin(pin):
        return redirect(url_for('employees', message='PIN must be 4 to 6 digits'))

    emp = {
        'name': name,
        'department': department,
        'position': position,
        'join_date': now_local().strftime('%Y-%m-%d'),
        'pin_hash': generate_password_hash(pin, method='pbkdf2:sha256'),
    }
    for _ in range(10):
        emp_id = f'EMP{next_counter():03d}'
        if store_create_employee(emp_id, emp):
            return redirect(url_for('employees', message=f'Employee {name} added successfully!'))
    return redirect(url_for('employees', message='Could not create the employee, please try again'))


@app.route('/set_pin', methods=['POST'])
@admin_required
def set_pin():
    emp_id = request.form.get('emp_id', '')
    pin = request.form.get('pin', '').strip()
    emp = store_get_employee(emp_id)
    if not emp:
        return redirect(url_for('employees', message='Employee not found'))
    if not valid_pin(pin):
        return redirect(url_for('employees', message='PIN must be 4 to 6 digits'))
    emp['pin_hash'] = generate_password_hash(pin, method='pbkdf2:sha256')
    emp['fails'] = 0
    emp['locked_until'] = 0
    store_set_employee(emp_id, emp)
    return redirect(url_for('employees', message=f'PIN updated for {emp["name"]}'))


@app.route('/delete_employee', methods=['POST'])
@admin_required
def delete_employee():
    emp_id = request.form.get('emp_id', '')
    if not store_get_employee(emp_id):
        return redirect(url_for('employees', message='Employee not found'))
    store_delete_employee(emp_id)
    return redirect(url_for('employees', message=f'Employee {emp_id} deleted'))


@app.route('/records')
@admin_required
def records():
    data = load_data()
    filter_date = request.args.get('date')
    records = list(data['attendance'])

    if filter_date:
        records = [r for r in records if r['date'] == filter_date]

    for record in records:
        if record.get('clock_out'):
            h = hours_between(record['date'], record.get('clock_in'), record['clock_out'])
            record['hours_worked'] = format_duration(h) if h is not None else 'Check times'
        else:
            record['hours_worked'] = 'In Progress'

    records = list(reversed(records))
    return render_template_string(RECORDS_TEMPLATE, records=records, filter_date=filter_date,
                                  current_month=now_local().strftime('%Y-%m'),
                                  message=request.args.get('message'),
                                  is_admin=session.get('is_admin', False))


def render_edit(emp_id, date, clock_in, clock_out, message=None):
    data = load_data()
    editing = bool(emp_id and valid_date(date) and store_get_record(date, emp_id))
    return render_template_string(EDIT_TEMPLATE, employees=data['employees'], editing=editing,
                                  emp_id=emp_id, date=date, clock_in=clock_in, clock_out=clock_out,
                                  message=message, is_admin=True)


@app.route('/edit_record')
@admin_required
def edit_record():
    emp_id = request.args.get('emp_id', '')
    date = request.args.get('date', '')
    rec = store_get_record(date, emp_id) if (emp_id and valid_date(date)) else None
    return render_edit(emp_id, date or now_local().strftime('%Y-%m-%d'),
                       (rec or {}).get('clock_in') or '', (rec or {}).get('clock_out') or '')


@app.route('/save_record', methods=['POST'])
@admin_required
def save_record():
    emp_id = request.form.get('emp_id', '')
    date = request.form.get('date', '')
    ci_raw = request.form.get('clock_in', '')
    co_raw = request.form.get('clock_out', '')
    emp = store_get_employee(emp_id) if emp_id else None
    ci = parse_time(ci_raw)
    co = parse_time(co_raw) if co_raw.strip() else None

    err = None
    if not emp:
        err = 'Please choose an employee.'
    elif not valid_date(date):
        err = 'Please enter a valid date.'
    elif not ci:
        err = 'Please enter a valid clock-in time.'
    elif co_raw.strip() and not co:
        err = 'Please enter a valid clock-out time.'
    elif co and hours_between(date, ci, co) is None:
        err = 'Clock-out must be later than clock-in.'
    if err:
        return render_edit(emp_id, date, ci_raw, co_raw, err)

    rec = store_get_record(date, emp_id) or {}
    rec.update({'employee_id': emp_id, 'employee_name': emp['name'], 'department': emp['department'],
                'date': date, 'clock_in': ci, 'clock_out': co})
    store_set_record(rec)
    return redirect(url_for('records', message=f'Record saved for {emp["name"]} on {date}'))


@app.route('/delete_record', methods=['POST'])
@admin_required
def delete_record():
    store_delete_record(request.form.get('date', ''), request.form.get('emp_id', ''))
    return redirect(url_for('records', message='Record deleted'))


@app.route('/reports')
@admin_required
def reports():
    data = load_data()
    week_ago_dt = now_local() - timedelta(days=7)
    week_ago = week_ago_dt.strftime('%Y-%m-%d')

    weekly_report = {}
    for record in data['attendance']:
        if record['date'] >= week_ago and record.get('clock_out'):
            emp_id = record['employee_id']
            if emp_id not in weekly_report:
                emp = data['employees'].get(emp_id, {})
                weekly_report[emp_id] = {
                    'name': emp.get('name', 'Unknown'),
                    'department': emp.get('department', 'Unknown'),
                    'days_present': 0,
                    'total_hours': 0.0
                }
            hours = hours_between(record['date'], record.get('clock_in'), record['clock_out']) or 0.0

            weekly_report[emp_id]['days_present'] += 1
            weekly_report[emp_id]['total_hours'] += hours

    for emp_id in weekly_report:
        days = weekly_report[emp_id]['days_present']
        weekly_report[emp_id]['avg_hours'] = weekly_report[emp_id]['total_hours'] / days if days > 0 else 0.0
        weekly_report[emp_id]['total_text'] = format_duration(weekly_report[emp_id]['total_hours'])
        weekly_report[emp_id]['avg_text'] = format_duration(weekly_report[emp_id]['avg_hours'])

    return render_template_string(REPORTS_TEMPLATE, weekly_report=weekly_report, is_admin=session.get('is_admin', False))

COMPANY_NAME = os.environ.get('COMPANY_NAME', 'Employee Attendance System')
DURATION_FMT = '[h]"h "mm"m"'
WORK_HOURS_PER_DAY = float(os.environ.get('WORK_HOURS_PER_DAY', '8'))


def build_export_rows(data, filter_date=None, month=None):
    rows = []
    for r in sorted(data['attendance'], key=lambda x: (x['date'], x.get('clock_in') or '')):
        if filter_date and r['date'] != filter_date:
            continue
        if month and not r['date'].startswith(month):
            continue
        if r.get('clock_out'):
            h = hours_between(r['date'], r.get('clock_in'), r['clock_out'])
            hours, status = (h, 'Completed') if h is not None else (None, 'Check times')
        else:
            hours, status = None, 'In Progress'
        try:
            day_name = datetime.strptime(r['date'], '%Y-%m-%d').strftime('%A')
        except Exception:
            day_name = ''
        rows.append({
            'date': r['date'], 'day': day_name, 'emp_id': r['employee_id'],
            'name': r.get('employee_name', ''), 'dept': r.get('department', ''),
            'clock_in': r.get('clock_in') or '', 'clock_out': r.get('clock_out') or '',
            'hours': hours, 'status': status,
        })
    return rows


def build_excel(rows, filter_date, month=None, report_title='Attendance Sheet'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    NAVY, BAND, WHITE = '1F3864', 'F2F5FA', 'FFFFFF'
    thin = Side(style='thin', color='BFC7D5')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name='Calibri', bold=True, color=WHITE, size=11)
    head_fill = PatternFill('solid', fgColor=NAVY)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center', indent=1)

    if month:
        period = datetime.strptime(month, '%Y-%m').strftime('%B %Y')
    elif filter_date:
        period = datetime.strptime(filter_date, '%Y-%m-%d').strftime('%d %B %Y')
    elif rows:
        d1 = datetime.strptime(rows[0]['date'], '%Y-%m-%d').strftime('%d %b %Y')
        d2 = datetime.strptime(rows[-1]['date'], '%Y-%m-%d').strftime('%d %b %Y')
        period = d1 if d1 == d2 else f'{d1} - {d2}'
    else:
        period = 'No records'
    generated = now_local().strftime('%d %B %Y, %I:%M %p')

    def title_block(ws, title, last_col):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
        ws['A1'] = COMPANY_NAME
        ws['A1'].font = Font(name='Calibri', bold=True, size=18, color=WHITE)
        for c in range(1, last_col + 1):
            ws.cell(row=1, column=c).fill = PatternFill('solid', fgColor=NAVY)
        ws['A2'].font = Font(name='Calibri', bold=True, size=14, color=NAVY)
        ws['A2'] = title
        ws['A2'].font = Font(name='Calibri', bold=True, size=13, color='444444')
        ws['A3'] = f'Period: {period}     |     Generated: {generated}'
        ws['A3'].font = Font(name='Calibri', italic=True, size=10, color='666666')
        ws.row_dimensions[1].height = 34
        for r in (1, 2, 3):
            ws.cell(row=r, column=1).alignment = left

    def header_row(ws, row, headers):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=i, value=h)
            c.font, c.fill, c.alignment, c.border = head_font, head_fill, center, border
        ws.row_dimensions[row].height = 22

    def page_setup(ws, header_row_num):
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f'{header_row_num}:{header_row_num}'
        ws.oddFooter.center.text = 'Page &P of &N'
        ws.sheet_view.showGridLines = False

    wb = Workbook()
    wb.properties.title = report_title
    wb.properties.creator = COMPANY_NAME

    # ---------------- Sheet 1: Daily attendance ----------------
    ws = wb.active
    ws.title = 'Attendance'
    ws.sheet_properties.tabColor = NAVY
    headers = ['Date', 'Day', 'Employee ID', 'Employee Name', 'Department',
               'Clock In', 'Clock Out', 'Time Worked', 'Status']
    title_block(ws, report_title, len(headers))
    HR = 5
    header_row(ws, HR, headers)

    status_style = {
        'Completed': ('1E7B34', 'E6F4EA'),
        'In Progress': ('9A6700', 'FFF4CE'),
        'Check times': ('B42318', 'FDE8E8'),
    }
    r = HR
    for i, row in enumerate(rows):
        r += 1
        vals = [
            datetime.strptime(row['date'], '%Y-%m-%d').date(), row['day'], row['emp_id'],
            row['name'], row['dept'],
            datetime.strptime(row['clock_in'], '%H:%M:%S').time() if row['clock_in'] else None,
            datetime.strptime(row['clock_out'], '%H:%M:%S').time() if row['clock_out'] else None,
            round(row['hours'] * 60) / 1440 if row['hours'] is not None else None, row['status'],
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = border
            c.font = Font(name='Calibri', size=11)
            c.alignment = left if col in (4, 5) else center
            if i % 2 == 1:
                c.fill = PatternFill('solid', fgColor=BAND)
        ws.cell(row=r, column=1).number_format = 'dd-mmm-yyyy'
        ws.cell(row=r, column=6).number_format = 'hh:mm AM/PM'
        ws.cell(row=r, column=7).number_format = 'hh:mm AM/PM'
        ws.cell(row=r, column=8).number_format = DURATION_FMT
        fg, bg = status_style[row['status']]
        sc = ws.cell(row=r, column=9)
        sc.font = Font(name='Calibri', bold=True, size=11, color=fg)
        sc.fill = PatternFill('solid', fgColor=bg)

    first, last = HR + 1, r
    if rows:
        r += 1
        ws.cell(row=r, column=1, value='TOTAL')
        ws.cell(row=r, column=8, value=f'=SUBTOTAL(109,H{first}:H{last})')
        ws.cell(row=r, column=9, value=f'=SUBTOTAL(103,A{first}:A{last})&" records"')
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=r, column=col)
            c.font = Font(name='Calibri', bold=True, size=11, color=NAVY)
            c.fill = PatternFill('solid', fgColor='DCE3F0')
            c.border = Border(top=Side(style='medium', color=NAVY), bottom=thin, left=thin, right=thin)
            c.alignment = center
        ws.cell(row=r, column=8).number_format = DURATION_FMT
        ws.auto_filter.ref = f'A{HR}:{get_column_letter(len(headers))}{last}'
    ws.freeze_panes = ws.cell(row=HR + 1, column=1)
    for i, w in enumerate([14, 13, 14, 24, 18, 12, 12, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    page_setup(ws, HR)

    # ---------------- Sheet 2: Summary per employee ----------------
    ws2 = wb.create_sheet('Summary')
    ws2.sheet_properties.tabColor = '2563EB'
    summary = {}
    for row in rows:
        e = summary.setdefault(row['emp_id'], {'name': row['name'], 'dept': row['dept'],
                                               'days': 0, 'hours': 0.0})
        if row['status'] == 'Completed':
            e['days'] += 1
            e['hours'] += row['hours']
    sh = ['Employee ID', 'Employee Name', 'Department', 'Days Present', 'Total Time',
          'Average / Day', f'Expected ({WORK_HOURS_PER_DAY:g}h/day)', 'Difference']
    title_block(ws2, report_title.replace('Attendance Sheet', 'Employee Summary'), len(sh))
    header_row(ws2, HR, sh)

    def summary_row(r, vals, diff_hours, fill=None, bold=False, color=None):
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=r, column=col, value=v)
            c.border = border
            c.font = Font(name='Calibri', size=11, bold=bold, color=color)
            c.alignment = left if col in (2, 3) else center
            if fill:
                c.fill = PatternFill('solid', fgColor=fill)
        d = ws2.cell(row=r, column=8)
        d.font = Font(name='Calibri', bold=True, size=11,
                      color='1E7B34' if diff_hours >= 0 else 'B42318')

    r = HR
    tot_days = tot_hours = 0
    for i, (eid, e) in enumerate(sorted(summary.items())):
        r += 1
        expected = e['days'] * WORK_HOURS_PER_DAY
        avg = e['hours'] / e['days'] if e['days'] else 0
        tot_days += e['days']
        tot_hours += e['hours']
        summary_row(r, [eid, e['name'], e['dept'], e['days'], format_duration(e['hours']),
                        format_duration(avg), format_duration(expected),
                        format_duration(e['hours'] - expected, signed=True)],
                    e['hours'] - expected, fill=BAND if i % 2 else None)
    if summary:
        r += 1
        tot_expected = tot_days * WORK_HOURS_PER_DAY
        summary_row(r, ['TOTAL', '', '', tot_days, format_duration(tot_hours),
                        format_duration(tot_hours / tot_days if tot_days else 0),
                        format_duration(tot_expected),
                        format_duration(tot_hours - tot_expected, signed=True)],
                    tot_hours - tot_expected, fill='DCE3F0', bold=True, color=NAVY)
        ws2.cell(row=r, column=8).font = Font(
            name='Calibri', bold=True, size=11,
            color='1E7B34' if tot_hours - tot_expected >= 0 else 'B42318')
    ws2.freeze_panes = ws2.cell(row=HR + 1, column=1)
    for i, w in enumerate([14, 24, 18, 14, 14, 16, 26, 14], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    page_setup(ws2, HR)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@app.route('/export')
@admin_required
def export_file():
    data = load_data()
    filter_date = request.args.get('date')
    fmt = request.args.get('format', 'xlsx')
    month = request.args.get('month')

    # Default: the current month. Use ?month=all for every record.
    if filter_date:
        month = None
    elif month == 'all':
        month = None
    elif not month or not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
        month = now_local().strftime('%Y-%m')

    rows = build_export_rows(data, filter_date, month)

    if month:
        month_label = datetime.strptime(month, '%Y-%m').strftime('%B %Y')
        report_title = f'{month_label} Attendance Sheet'
    elif filter_date:
        day_label = datetime.strptime(filter_date, '%Y-%m-%d').strftime('%d %B %Y')
        report_title = f'Attendance Sheet - {day_label}'
    else:
        report_title = 'Attendance Sheet - All Records'
    base_name = report_title

    if fmt != 'csv':
        try:
            output = build_excel(rows, filter_date, month, report_title)
            return send_file(
                output,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                as_attachment=True, download_name=base_name + '.xlsx')
        except ImportError:
            pass  # openpyxl not installed -> fall back to CSV below

    si = io.StringIO()
    writer = csv.writer(si)
    writer.writerow(['Date', 'Day', 'Employee ID', 'Employee Name', 'Department',
                     'Clock In', 'Clock Out', 'Time Worked', 'Status'])
    for x in rows:
        writer.writerow([x['date'], x['day'], x['emp_id'], x['name'], x['dept'],
                         x['clock_in'], x['clock_out'],
                         format_duration(x['hours']), x['status']])
    output = io.BytesIO()
    output.write(b'\xef\xbb\xbf')  # BOM so Excel reads the file correctly
    output.write(si.getvalue().encode('utf-8'))
    output.seek(0)
    return send_file(output, mimetype='text/csv', as_attachment=True,
                     download_name=base_name + '.csv')


@app.route('/health')
def health():
    info = {'redis_configured': bool(REDIS_URL and REDIS_TOKEN)}
    if REDIS_URL and REDIS_TOKEN:
        try:
            info['redis_ping'] = _cmd('PING')
        except Exception as e:
            info['redis_error'] = str(e)
    return jsonify(info)


if __name__ == '__main__':
    app.run(debug=True, port=5000)
